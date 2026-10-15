"""library.py - reads the owner's spreadsheets (strategies + bots) and keeps a JSON snapshot so the app also
runs without openpyxl. The spreadsheets stay the source of truth: edit them, the snapshot is rebuilt."""
from __future__ import annotations
import json, os
from . import ROOT
from .strategies import NEW_ID_MAP, OVERLAYS, GENERATORS

DATA = os.path.join(ROOT, "data")
STRAT_XLSX = os.path.join(DATA, "JARVUS_2_STRATEGIES.xlsx")
BOTS_XLSX = os.path.join(DATA, "JARVUS_3_BOTS.xlsx")
SNAPSHOT = os.path.join(DATA, "library.json")

KEEP_KEYS = ["tier", "starter", "id", "job", "jarvus_grade", "catalog_grade", "evidence", "caveat", "risk", "bars",
             "markets", "entry", "exit_rule", "stop", "settings", "crypto_override", "btc_test"]
BOT_KEYS = {"Agent": "bot_id", "Status": "status", "Priority": "priority", "Group": "group", "Role": "role",
            "Instrument": "symbol", "Bar": "bar", "Old strategy": "old_strategy", "Old test Sharpe": "old_sharpe",
            "Old test trades": "old_trades", "Old DSR": "old_dsr", "Old PBO": "old_pbo", "Old verdict": "verdict",
            "New candidate 1": "candidate_1", "New candidate 2": "candidate_2", "New candidate 3": "candidate_3",
            "Day trade settings": "day_trade_settings", "Canada venue check": "venue_check", "Why": "why",
            "Training data": "training_data", "Allowed to trade": "allowed", "Min stop (ATR)": "min_stop_atr",
            "Weekend entries": "weekend_entries", "Stale if bar older than": "stale_rule"}


def _rows(ws):
    return [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]


def _build_from_xlsx() -> dict:
    import openpyxl
    s = openpyxl.load_workbook(STRAT_XLSX, read_only=True, data_only=True)
    keep = []
    for r in _rows(s["Keep"])[1:]:
        d = dict(zip(KEEP_KEYS, r))
        g, f, e = d["id"].split(".")
        d.update(generator=g, filter=f, exit=e, runnable=g in GENERATORS, starter=d["starter"] == "Yes")
        keep.append(d)
    new = []
    for r in _rows(s["New (to build)"])[1:]:
        d = dict(zip(["id", "family", "job", "evidence", "card", "rules", "bars", "status"], r))
        d["runnable"] = d["id"] in NEW_ID_MAP
        d["overlay"] = d["id"] in OVERLAYS
        d["built"] = d["runnable"] or d["overlay"]
        new.append(d)
    def table(name):
        rows = _rows(s[name]); return {"header": list(rows[0]), "rows": [list(r) for r in rows[1:]]}
    b = openpyxl.load_workbook(BOTS_XLSX, read_only=True, data_only=True)
    rows = _rows(b["Bots"]); hdr = rows[0]
    bots = [{BOT_KEYS.get(h, h): v for h, v in zip(hdr, r)} for r in rows[1:] if r[0]]
    roles_rows = _rows(b["Roles"])
    roles = [dict(zip(["role", "bar", "bots", "candidate_1", "candidate_2", "candidate_3", "old_trades", "plan", "extra"], r))
             for r in roles_rows[1:] if r[1]]
    return {"keep": keep, "new": new, "do_not_build": table("Do not build")["rows"], "cut": table("Cut"),
            "day_trade_rules": table("Day-trade rules"), "cost_in_r": [list(r) for r in _rows(s["Cost in R"])],
            "measured": table("Measured (BTC)"), "bots": bots, "roles": roles,
            "auto_pause_rules": table_b(b, "Auto-pause rules"), "fleet_feed": table_b(b, "Fleet feed"),
            "next_steps": table_b(b, "Next steps")}


def table_b(wb, name):
    rows = _rows(wb[name]); return {"header": list(rows[0]), "rows": [list(r) for r in rows[1:]]}


_CACHE: dict = {}


def load(force: bool = False) -> dict:
    if _CACHE and not force: return _CACHE
    data = None
    xlsx_newer = os.path.exists(STRAT_XLSX) and (not os.path.exists(SNAPSHOT) or
                 max(os.path.getmtime(STRAT_XLSX), os.path.getmtime(BOTS_XLSX)) > os.path.getmtime(SNAPSHOT))
    if xlsx_newer or force:
        try:
            data = _build_from_xlsx()
            with open(SNAPSHOT, "w", encoding="utf-8") as f: json.dump(data, f, indent=1, default=str)
        except Exception:
            data = None
    if data is None:
        with open(SNAPSHOT, encoding="utf-8") as f: data = json.load(f)
    _CACHE.clear(); _CACHE.update(data)
    return _CACHE


def strategy(sid: str) -> dict | None:
    lib = load()
    for d in lib["keep"] + lib["new"]:
        if d["id"] == sid: return d
    return None


def bots() -> list: return load()["bots"]
def roles() -> list: return load()["roles"]


def bot(bot_id: str) -> dict | None:
    return next((b for b in bots() if b["bot_id"] == bot_id), None)


def is_meme(b: dict) -> bool:
    return "meme" in (b.get("role") or "").lower() or b.get("group") == "Meme coins"


def settings_for(sid: str) -> str:
    d = strategy(sid)
    return (d or {}).get("settings") or ""
